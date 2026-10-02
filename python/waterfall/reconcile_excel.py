"""Recalculates the workbook in headless LibreOffice for several cases and compares to the Python model.

Run: py -3 -m waterfall.reconcile_excel   (needs LibreOffice installed)
"""
import json
import shutil
import subprocess
import sys
import tempfile
from dataclasses import replace
from pathlib import Path
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter as L
from . import config as cfg, metrics as mt, waterfall as wf
from .waterfall import col

# ------------------------------------------------------------------ USER INPUTS
XLSX = Path(__file__).resolve().parents[2] / "excel" / "securitization_waterfall.xlsx"
META = XLSX.with_suffix(".meta.json")
REPORT = Path(__file__).resolve().parents[1] / "outputs" / "excel_reconciliation.json"
SOFFICE_CANDIDATES = [r"C:\Program Files\LibreOffice\program\soffice.exe",
                      r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]
TOL_CASH = 0.01          # USD per cash flow cell
TOL_RATIO = 1e-6         # WAL years and annualised yields
CASES_EXTRA = [          # (label, scenario, prorata, hedge) beyond the six standard scenarios
    ("pro rata Base", cfg.BASE, 1, 0), ("pro rata Stress", cfg.STRESS, 1, 0),
    ("pro rata Rate shock", cfg.SCENARIOS[5], 1, 0), ("hedged Rate shock", cfg.SCENARIOS[5], 0, 1),
    ("arrears sweep A", cfg.Scenario("custom", 0.02, 0.18, 0.20, 12), 0, 0),
    ("arrears sweep B", cfg.Scenario("custom", 0.10, 0.40, 0.70, 6), 0, 0),
]
# ------------------------------------------------------------------------------


def soffice():
    for p in SOFFICE_CANDIDATES:
        if Path(p).exists():
            return p
    sys.exit("LibreOffice not found; install it (winget install TheDocumentFoundation.LibreOffice)")


def recalc(src, scn, prorata, hedge, slot, meta, work):
    wb = openpyxl.load_workbook(src)
    ws = wb["Inputs"]
    row = meta["scn_first"] + slot - 1
    for j, v in enumerate([scn.name, scn.cpr, scn.cdr, scn.severity, scn.lag, scn.index_shift], start=1):
        ws.cell(row, j, v)
    names = meta["names"]
    cell = lambda key: names[key].split("!")[1].replace("$", "")
    ws[cell("sel")] = slot
    ws[cell("prorata")] = prorata
    ws[cell("hedge")] = hedge
    tmp = work / "case.xlsx"
    wb.save(tmp)
    out = work / "out"
    shutil.rmtree(out, ignore_errors=True)
    subprocess.run([soffice(), "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(tmp)],
                   check=True, capture_output=True, timeout=600)
    return openpyxl.load_workbook(out / "case.xlsx", data_only=True)


def series(wsx, colidx, meta):
    return np.array([wsx.cell(r, colidx).value or 0.0 for r in range(meta["first"], meta["last"] + 1)], float)


def compare(label, scn, prorata, hedge, wbv, meta):
    deal = cfg.Deal(pro_rata_start=bool(prorata), hedged=bool(hedge))
    df = wf.results_frame(deal, scn)
    w = wbv["Waterfall"]
    c = meta["wf"]
    worst_cash, worst_ratio, issues = 0.0, 0.0, []
    pairs = []
    for k in cfg.CLASSES:
        pairs += [(f"bal_{k}", df[col("bal", k)].values), (f"int_{k}", df[col("int_paid", k)].values),
                  (f"prin_{k}", df[col("prin_paid", k)].values), (f"wd_{k}", df[col("writedown", k)].values),
                  (f"carry_{k}", df[col("unpaid_int", k)].values)]
    pairs += [("resid_cash", df[wf.RESID_CASH].values), ("swept", df[wf.ARREARS_SWEPT].values),
              ("reserve_end", df[wf.RES_END].values), ("fees_paid", df[wf.FEES_PAID].values),
              ("draw", df[wf.RES_DRAW].values), ("excess_spread", df[wf.EXCESS_SPREAD].values)]
    for name, py in pairs:
        xl = series(w, c[name], meta)
        d = float(np.abs(xl - py).max())
        worst_cash = max(worst_cash, d)
        if d > TOL_CASH:
            issues.append(f"{name}: max diff {d:.4f}")
    summ = mt.summarize(df, deal, scn)
    s = wbv["Scenario_Summary"]
    for i, k in enumerate(cfg.CLASSES + ["R"]):
        r = 4 + i
        for header, colno, key in (("WAL", 3, "wal"), ("yield", 4, "yield_at_par")):
            xl, py = s.cell(r, colno).value, summ[k][key]
            if py is None or xl in (None, "n/a"):
                if not (py is None and xl in (None, "n/a")):
                    issues.append(f"{k} {header}: excel {xl} python {py}")
                continue
            if isinstance(xl, str):
                issues.append(f"{k} {header}: excel returned {xl}, python {py}")
                continue
            worst_ratio = max(worst_ratio, abs(xl - py))
            if abs(xl - py) > TOL_RATIO:
                issues.append(f"{k} {header}: excel {xl} python {py}")
        xl_wd = s.cell(r, 9).value
        if abs(xl_wd - summ[k]["writedown"]) > TOL_CASH:
            issues.append(f"{k} write-down: excel {xl_wd} python {summ[k]['writedown']}")
        if bool(s.cell(r, 12).value) != summ[k]["any_shortfall"]:
            issues.append(f"{k} shortfall flag differs")
    lab = mt.class_labels(summ)
    if s["B10"].value != lab["junior_writedown"] or s["B11"].value != lab["senior_writedown"]:
        issues.append(f"labels: excel {s['B10'].value}/{s['B11'].value} python {lab['junior_writedown']}/{lab['senior_writedown']}")
    months = int(df[wf.BREACH_ANY].sum())
    if s["B17"].value != months:
        issues.append(f"breach months: excel {s['B17'].value} python {months}")
    if s["B19"].value != 1:
        issues.append("workbook own checks failed")
    return dict(case=label, passed=not issues, worst_cash_diff=worst_cash, worst_ratio_diff=worst_ratio, issues=issues)


def main():
    meta = json.loads(META.read_text())
    work = Path(tempfile.mkdtemp())
    cases = [(s.name, s, 0, 0) for s in cfg.SCENARIOS] + CASES_EXTRA
    results = []
    for i, (label, scn, prorata, hedge) in enumerate(cases):
        slot = cfg.SCENARIOS.index(scn) + 1 if scn in cfg.SCENARIOS and label == scn.name else 7
        wbv = recalc(XLSX, scn, prorata, hedge, slot, meta, work)
        res = compare(label, scn, prorata, hedge, wbv, meta)
        results.append(res)
        print(("PASS" if res["passed"] else "FAIL"), label, f"cash {res['worst_cash_diff']:.2e}", res["issues"][:3])
    REPORT.write_text(json.dumps(dict(all_passed=all(r["passed"] for r in results), cases=results), indent=1))
    print("all passed:", all(r["passed"] for r in results))


if __name__ == "__main__":
    main()
